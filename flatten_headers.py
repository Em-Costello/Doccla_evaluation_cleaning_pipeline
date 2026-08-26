#!/usr/bin/env python3
"""
flatten_headers.py

Flattens a 2-row merged-cell Excel header (a "group" row over a "sub-field"
row) into a single row of clean column names, and writes the result out as
a CSV.

Logic:
  1. Row 1 holds group labels (e.g. "COPD Diagnosis"), often merged across
     several columns. We forward-fill each merged group label across all
     the columns it spans.
  2. Row 2 holds the sub-field labels (e.g. "Code Term", "Date"). Any
     column whose row-2 value is blank is treated as an empty/unused
     column (like column H in the source file) and is dropped entirely.
  3. The flattened header for a kept column is "<Group> - <Subfield>",
     e.g. "COPD Diagnosis - Code Term". Groups listed in
     NO_PREFIX_GROUPS are exempt from the prefix, since qualifying them
     adds noise (e.g. "Patient Details - NHS Number" -> just "NHS Number").
  4. Data rows start after the header rows (default: row 3 onward) and are
     written out to CSV using only the kept columns.

Usage:
    python flatten_headers.py input.xlsx output.csv
    python flatten_headers.py input.xlsx output.csv --sheet "Sheet1"
    python flatten_headers.py input.xlsx output.csv --header-rows 2
"""

import argparse
import csv
import sys

try:
    import openpyxl
except ImportError:
    sys.exit(
        "This script needs the 'openpyxl' package.\n"
        "Install it with:  pip install openpyxl"
    )

# Group names that should NOT be prefixed onto their sub-field name.
# Add to this set if other groups should also be left bare.
NO_PREFIX_GROUPS = {"Patient Details"}


def find_header_start_row(ws, max_col, max_scan_rows=10):
    """
    Returns the first row (1-indexed) within the first `max_scan_rows`
    rows that isn't entirely blank. Handles files that have a blank
    leading row (or rows) before the real header.
    """
    for r in range(1, max_scan_rows + 1):
        for c in range(1, max_col + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                return r
    return 1  # fallback


def forward_fill_row(ws, row_idx, max_col):
    """
    Return a list (1-indexed via offset) of values for `row_idx`, with
    merged-cell values propagated across every column they span.
    """
    values = [ws.cell(row=row_idx, column=c).value for c in range(1, max_col + 1)]

    for merge_range in ws.merged_cells.ranges:
        if merge_range.min_row <= row_idx <= merge_range.max_row:
            anchor_value = ws.cell(row=merge_range.min_row, column=merge_range.min_col).value
            for c in range(merge_range.min_col, merge_range.max_col + 1):
                values[c - 1] = anchor_value

    return values


def build_flat_headers(ws, header_rows=2, start_row=None):
    """
    Returns (kept_column_indices, flat_headers, data_start_row).
    kept_column_indices are 1-indexed column numbers to pull from each
    data row; flat_headers is the corresponding list of flattened header
    strings; data_start_row is the first row (1-indexed) after the header
    block.
    """
    max_col = ws.max_column

    if start_row is None:
        start_row = find_header_start_row(ws, max_col)

    if header_rows == 1:
        headers = forward_fill_row(ws, start_row, max_col)
        kept_cols, flat_headers = [], []
        for c, h in enumerate(headers, start=1):
            if h is None or str(h).strip() == "":
                continue
            kept_cols.append(c)
            flat_headers.append(str(h).strip())
        return kept_cols, flat_headers, start_row + header_rows

    # Two-row header case (the COPD dataset structure)
    group_row = forward_fill_row(ws, start_row, max_col)
    sub_row_raw = [ws.cell(row=start_row + 1, column=c).value for c in range(1, max_col + 1)]

    kept_cols = []
    flat_headers = []
    for c in range(1, max_col + 1):
        sub_val = sub_row_raw[c - 1]
        if sub_val is None or str(sub_val).strip() == "":
            # Blank sub-header -> unused/empty column (e.g. column H). Drop it.
            continue

        sub_val = str(sub_val).strip()
        group_val = group_row[c - 1]
        group_val = str(group_val).strip() if group_val else ""

        if not group_val or group_val in NO_PREFIX_GROUPS:
            flat_header = sub_val
        else:
            flat_header = f"{group_val} - {sub_val}"

        kept_cols.append(c)
        flat_headers.append(flat_header)

    return kept_cols, flat_headers, start_row + header_rows


def main():
    parser = argparse.ArgumentParser(description="Flatten a 2-row merged Excel header into a single-row CSV.")
    parser.add_argument("input_xlsx", help="Path to the source .xlsx file")
    parser.add_argument("output_csv", help="Path to write the flattened .csv file")
    parser.add_argument("--sheet", default=None, help="Sheet name (default: active sheet)")
    parser.add_argument("--header-rows", type=int, default=2, choices=[1, 2],
                         help="Number of header rows in the source file (default: 2)")
    parser.add_argument("--start-row", type=int, default=None,
                         help="Row number (1-indexed) where the header block starts. "
                              "Default: auto-detected as the first non-blank row.")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.input_xlsx, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active

    kept_cols, flat_headers, data_start_row = build_flat_headers(
        ws, header_rows=args.header_rows, start_row=args.start_row
    )

    print(f"Header block detected starting at row {data_start_row - args.header_rows}; "
          f"data starts at row {data_start_row}.")
    print(f"Kept {len(kept_cols)} of {ws.max_column} columns.")
    dropped = [c for c in range(1, ws.max_column + 1) if c not in kept_cols]
    if dropped:
        dropped_letters = [openpyxl.utils.get_column_letter(c) for c in dropped]
        print(f"Dropped empty/unused columns: {', '.join(dropped_letters)}")

    with open(args.output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(flat_headers)

        row_count = 0
        for row in ws.iter_rows(min_row=data_start_row, values_only=True):
            row_values = [row[c - 1] if row[c - 1] is not None else "" for c in kept_cols]
            # Skip fully blank rows
            if any(str(v).strip() for v in row_values):
                writer.writerow(row_values)
                row_count += 1

    print(f"Wrote {row_count} data row(s) to {args.output_csv}")


if __name__ == "__main__":
    main()