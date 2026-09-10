import pandas as pd
import numpy as np
import re
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NEW_EXPECTED_COLUMNS = [ # CHANGE IF THE SOURCE FILE STRUCTURE CHANGES
    'NHS Number', 'Date of Birth', 'Age', 'Gender', 'Postcode',
    'Ethnic Origin', 'Clinical Codes - Code Term', 'Clinical Codes - Date',
    'COPD Diagnosis - Code Term', 'COPD Diagnosis - Date',
    'Depression - Code Term', 'Depression - Date', 'Anxiety - Code Term',
    'Anxiety - Date', 'Referral to Pulmonary Rehab - Code Term',
    'Referral to Pulmonary Rehab - Date', 'Smoking Status - Code Term',
    'Smoking Status - Date', 'Smoking Status - Associated Text',
    'Review of Inhaler Technique - Code Term',
    'Review of Inhaler Technique - Date',
    'Medication Issues - Name, Dosage and Quantity',
    'Medication Issues - Date of Issue', 'Medication Issues - Dose',
    'Prednisolone - Name, Dosage and Quantity',
    'Prednisolone - Date of Issue', 'Prednisolone - Dose',
    'Antibiotics - Name, Dosage and Quantity',
    'Antibiotics - Date of Issue', 'Antibiotics - Dose',
    'Consultations - Date', 'Consultations - Type of Consultation',
    'Cardiovascular Disease / Heart Disease - Code Term',
    'Cardiovascular Disease / Heart Disease - Date',
    'Cardiovascular Disease / Heart Disease - Associated Text',
    'Influenza - Code Term', 'Influenza - Date', 'Pneumococcal - Code Term',
    'Pneumococcal - Date', 'COVID - Code Term', 'COVID - Date'
]

RENAME_MAP = {
    'NHS Number': 'NHS_Number',
    'Date of Birth': 'DOB',
    'Ethnic Origin': 'Ethnicity',
    'Clinical Codes - Code Term': 'MRC_Code',
    'Clinical Codes - Date': 'MRC_Code_Date',
    'COPD Diagnosis - Code Term': 'COPD_Diagnosis',
    'COPD Diagnosis - Date': 'COPD_Diagnosis_Date',
    'Depression - Code Term': 'Depression',
    'Depression - Date': 'Depression_Date',
    'Anxiety - Code Term': 'Anxiety',
    'Anxiety - Date': 'Anxiety_Date',
    'Referral to Pulmonary Rehab - Code Term': 'Referral_to_Pulmonary_Rehab',
    'Referral to Pulmonary Rehab - Date': 'Referral_to_Pulmonary_Rehab_Date',
    'Smoking Status - Code Term': 'Smoking_Status',
    'Smoking Status - Date': 'Smoking_Status_Date',
    'Smoking Status - Associated Text': 'Smoking_Status_Associated_Text',
    'Review of Inhaler Technique - Code Term': 'Review_of_Inhaler_Technique',
    'Review of Inhaler Technique - Date': 'Review_of_Inhaler_Technique_Date',
    'Medication Issues - Name, Dosage and Quantity': 'Inhaler_Name_Dosage_Quantity',
    'Medication Issues - Date of Issue': 'Inhaler_Issue_Date',
    'Medication Issues - Dose': 'Inhaler_Dose',
    'Prednisolone - Name, Dosage and Quantity': 'Prednisolone_Name_Dosage_Quantity',
    'Prednisolone - Date of Issue': 'Prednisolone_Issue_Date',
    'Prednisolone - Dose': 'Prednisolone_Dose',
    'Antibiotics - Name, Dosage and Quantity': 'Antibiotic_Name_Dosage_Quantity',
    'Antibiotics - Date of Issue': 'Antibiotic_Issue_Date',
    'Antibiotics - Dose': 'Antibiotic_Dose',
    'Consultations - Date': 'Consultation_Date',
    'Consultations - Type of Consultation': 'Consultation_Type',
    'Cardiovascular Disease / Heart Disease - Code Term': 'Cardiovascular_Disease',
    'Cardiovascular Disease / Heart Disease - Date': 'Cardiovascular_Disease_Date',
    'Cardiovascular Disease / Heart Disease - Associated Text': 'Cardiovascular_Disease_Associated_Text',
    'Influenza - Code Term': 'Influenza_Vaccine',
    'Influenza - Date': 'Influenza_Vaccine_Date',
    'Pneumococcal - Code Term': 'Pneumococcal_Vaccine',
    'Pneumococcal - Date': 'Pneumococcal_Vaccine_Date',
    'COVID - Code Term': 'COVID_Vaccine',
    'COVID - Date': 'COVID_Vaccine_Date',
}

PATIENT_LEVEL_COLS = ['MRC_Code', 'COPD_Diagnosis', 'Depression', 'Anxiety',
                       'Smoking_Status', 'Gender', 'Ethnicity', 'Cardiovascular_Disease']

PATIENT_COLS = ['Patient_ID', 'GP_ID', 'Postcode', 'Age', 'DOB', 'Gender', 'Ethnicity',
                 'MRC_Code', 'MRC_Code_Date', 'COPD_Diagnosis', 'COPD_Diagnosis_Date',
                 'Depression', 'Depression_Date', 'Anxiety', 'Anxiety_Date',
                 'Smoking_Status', 'Smoking_Status_Date',
                 'Cardiovascular_Disease', 'Cardiovascular_Disease_Date']

RESCUE_PACK_TEXT_COLS = ['Inhaler_Dose', 'Prednisolone_Dose', 'Antibiotic_Dose']

# ---------- GP sources to process ----------
GP_SOURCES = [
    {'gp_id': 'GP01', 'source_file': 'evaluation_data/flattened/WHC_GP01.csv'},
    {'gp_id': 'GP02', 'source_file': 'evaluation_data/flattened/HARP_GP02.csv'},
    {'gp_id': 'GP03', 'source_file': 'evaluation_data/flattened/LW_GP03.csv'},
    {'gp_id': 'GP04', 'source_file': 'evaluation_data/flattened/MGP_GP04.csv'},
    {'gp_id': 'GP05', 'source_file': 'evaluation_data/flattened/SLRP_GP05.csv'},
    {'gp_id': 'GP06', 'source_file': 'evaluation_data/flattened/VS_GP06.csv'},
    
]

# ---------- Load categorical mappings once (shared across all GP files) ----------
diagnosis_mapping = pd.read_csv(Path('mapping_files/copd_diagnosis_mapping.csv')).set_index('raw_text_lowercase')
COPD_DIAGNOSIS_CATEGORY_MAP = diagnosis_mapping['category'].to_dict()

gender_mapping = pd.read_csv(Path('mapping_files/gender_mapping.csv')).set_index('raw_text_lowercase')
GENDER_CATEGORY_MAP = gender_mapping['category'].to_dict()

ethnicity_mapping = pd.read_csv(Path('mapping_files/ethnicity_mapping.csv')).set_index('raw_text_lowercase')
ETHNICITY_CATEGORY_MAP = ethnicity_mapping['category'].to_dict()

cardiovascular_mapping = pd.read_csv(Path('mapping_files/cardiovascular_disease_mapping.csv')).set_index('raw_text_lowercase')
CARDIOVASCULAR_DISEASE_CATEGORY_MAP = cardiovascular_mapping['category'].to_dict()


def classify_smoking_status(raw):
    if pd.isna(raw):
        return pd.NA
    text = str(raw).strip().lower()
    if re.search(r'non[- ]smoker|never smoked', text):
        return 'Non-Smoker'

    has_smoker_word = 'smoker' in text
    has_rolls_own = 'cigarette' in text and 'roll' in text
    has_cigar = 'cigar' in text and 'cigarette' not in text
    has_consumption_or_packyears = 'cigarette' in text and ('consum' in text or 'pack' in text and 'year' in text)

    if not (has_smoker_word or has_rolls_own or has_cigar or has_consumption_or_packyears):
        return None

    status = 'Former' if re.search(r'\bex[- ]', text) else 'Current'
    if 'very heavy' in text:
        heaviness, band = 'Very Heavy', '40+/day'
    elif 'heavy' in text:
        heaviness, band = 'Heavy', '20-39/day'
    elif 'moderate' in text:
        heaviness, band = 'Moderate', '10-19/day'
    elif 'light' in text:
        heaviness, band = 'Light', '1-9/day'
    elif 'trivial' in text:
        heaviness, band = 'Trivial', '<1/day'
    else:
        return f'{status} Smoker (Amount Not Specified)'
    return f'{status} {heaviness} Smoker ({band})'


def is_valid_nhs_number(nhs_num):
    digits = re.sub(r'\D', '', str(nhs_num))
    if len(digits) != 10:
        return False
    total = sum(int(d) * w for d, w in zip(digits[:9], range(10, 1, -1)))
    check_digit = 11 - (total % 11)
    if check_digit == 11:
        check_digit = 0
    if check_digit == 10:
        return False
    return check_digit == int(digits[9])


def build_event_table(source_df, cols, subset_for_filter=None):
    subset_for_filter = subset_for_filter or cols
    mask = source_df[subset_for_filter].notna().any(axis=1)
    return source_df.loc[mask, ['Patient_ID', 'GP_ID'] + cols].reset_index(drop=True)


def process_gp_file(gp_id, source_file):
    """Run the full single-file cleaning pipeline for one GP surgery extract
    and return a dict of {table_name: DataFrame}, each tagged with GP_ID.
    Every error message is prefixed with the gp_id so failures are traceable
    back to the offending file when processing several GPs in one run."""

    def fail(msg):
        raise ValueError(f"[{gp_id}] {msg}")

    # ---------- 1. Load + validate structure ----------
    df = pd.read_csv(source_file)
    df.columns = df.columns.str.strip()
    if list(df.columns) != NEW_EXPECTED_COLUMNS:
        fail("Column mismatch - structure of this file differs from the expected format.")

    # ---------- 2. Rename to internal names ----------
    df.rename(columns=RENAME_MAP, inplace=True)
    df.drop(columns=['Review_of_Inhaler_Technique', 'Review_of_Inhaler_Technique_Date'], inplace=True)

    # ---------- 2.5 Clean inconsistent data formats ----------
    diagnosis_normalized = df['COPD_Diagnosis'].astype(str).str.replace(',', '', regex=False).str.strip().str.lower().replace('nan', pd.NA)
    unmapped = set(diagnosis_normalized.dropna().unique()) - set(COPD_DIAGNOSIS_CATEGORY_MAP.keys())
    if unmapped:
        fail(f"Unmapped COPD_Diagnosis value(s) - add rows to copd_diagnosis_mapping.csv: {unmapped}")
    df['COPD_Diagnosis'] = diagnosis_normalized.map(COPD_DIAGNOSIS_CATEGORY_MAP)

    df['MRC_Code'] = pd.to_numeric(df['MRC_Code'].str.extract(r'(\d+)')[0], errors='coerce')
    if not df['MRC_Code'].dropna().isin([1, 2, 3, 4, 5]).all():
        fail("MRC_Code contains unexpected values - check file structure.")

    gender_normalized = df['Gender'].astype(str).str.replace(',', '', regex=False).str.strip().str.lower().replace('nan', pd.NA)
    unmapped_gender = set(gender_normalized.dropna().unique()) - set(GENDER_CATEGORY_MAP.keys())
    if unmapped_gender:
        fail(f"Unmapped Gender value(s) - add rows to gender_mapping.csv: {unmapped_gender}")
    df['Gender'] = gender_normalized.map(GENDER_CATEGORY_MAP)

    cardiovascular_normalized = df['Cardiovascular_Disease'].astype(str).str.replace(',', '', regex=False).str.strip().str.lower().replace('nan', pd.NA)
    unmapped_cardiovascular = set(cardiovascular_normalized.dropna().unique()) - set(CARDIOVASCULAR_DISEASE_CATEGORY_MAP.keys())
    if unmapped_cardiovascular:
        fail(f"Unmapped Cardiovascular_Disease value(s) - add rows to cardiovascular_disease_mapping.csv: {unmapped_cardiovascular}")
    df['Cardiovascular_Disease'] = cardiovascular_normalized.map(CARDIOVASCULAR_DISEASE_CATEGORY_MAP)

    ethnicity_normalized = (
            df['Ethnicity']
            .astype(str)
            .str.replace(',', '', regex=False)
            .str.strip()
            .str.lower()
            .replace('nan', pd.NA)
        )
    unmapped_ethnicity = set(ethnicity_normalized.dropna().unique()) - set(ETHNICITY_CATEGORY_MAP.keys())
    if unmapped_ethnicity:
        fail(
                f"Unmapped Ethnicity value(s) - add rows to ethnicity_mapping.csv "
                f"(check the ONS SNOMED-to-harmonised-category crosswalk first): {unmapped_ethnicity}"
            )
    df['Ethnicity'] = ethnicity_normalized.map(ETHNICITY_CATEGORY_MAP)

    smoking_raw = df['Smoking_Status']
    smoking_classified = smoking_raw.apply(classify_smoking_status)
    unrecognised_mask = smoking_raw.notna() & smoking_classified.isna()
    if unrecognised_mask.any():
        fail(
            f"Unrecognised Smoking_Status pattern(s) - extend classify_smoking_status(): "
            f"{sorted(smoking_raw[unrecognised_mask].unique())}"
        )
    df['Smoking_Status'] = smoking_classified

    # ---------- 3. Patient ID: real NHS Number (confirmed patient-row only) ----------
    is_anchor = df['Age'].notna() & df['Postcode'].notna() & df['DOB'].notna()
    if not is_anchor.iloc[0]:
        fail("First row isn't a patient anchor row - check file structure.")

    nhs_number_mismatch = (df['NHS_Number'].notna() != is_anchor).sum()
    if nhs_number_mismatch:
        fail(
            f"{nhs_number_mismatch} row(s) where NHS_Number presence doesn't match the "
            f"anchor-row pattern - re-run diagnostics before proceeding."
        )

    anchor_rows = df.loc[is_anchor].copy()
    nhs_valid = anchor_rows['NHS_Number'].apply(is_valid_nhs_number)
    if not nhs_valid.all():
        n_invalid = (~nhs_valid).sum()
        invalid_row_positions = anchor_rows.index[~nhs_valid].tolist()
        fail(
            f"{n_invalid} NHS Number(s) failed the checksum validation, at row positions "
            f"{invalid_row_positions} (not printing the numbers themselves - check those "
            f"rows directly in your own environment before proceeding)."
        )

    dup_nhs = anchor_rows['NHS_Number'][anchor_rows['NHS_Number'].duplicated(keep=False)]
    if not dup_nhs.empty:
        fail(
            f"Duplicate NHS Number(s) found across separate patient blocks within this "
            f"file - review manually: {sorted(dup_nhs.unique())}"
        )

    temp_group = is_anchor.cumsum()
    group_to_id = dict(zip(temp_group[is_anchor], anchor_rows['NHS_Number']))
    df['Patient_ID'] = temp_group.map(group_to_id)
    df['GP_ID'] = gp_id

    off_anchor_leak = df.loc[~is_anchor, PATIENT_LEVEL_COLS].notna().any().any()
    if off_anchor_leak:
        fail("A patient-level field is filled on a non-anchor row - "
             "re-run diagnostics, this file may not match the assumed structure.")

    # ---------- 3.5 Parse dates needed for rescue pack matching (ahead of section 6) ----------
    RESCUE_PACK_DATE_COLS = ['Inhaler_Issue_Date', 'Prednisolone_Issue_Date', 'Antibiotic_Issue_Date']
    for c in RESCUE_PACK_DATE_COLS:
        df[c] = pd.to_datetime(df[c], dayfirst=True, errors='coerce')

    # ---------- 4. patients_df: one row per patient ----------
    patients_df = df.loc[is_anchor, PATIENT_COLS].reset_index(drop=True)
    patients_df['Depression'] = patients_df['Depression'].notna().map({True: 'Yes', False: pd.NA})
    patients_df['Anxiety'] = patients_df['Anxiety'].notna().map({True: 'Yes', False: pd.NA})

    # ---------- 5. Event tables ----------
    inhaler_prescriptions_df = build_event_table(
        df, ['Inhaler_Name_Dosage_Quantity', 'Inhaler_Dose', 'Inhaler_Issue_Date'])
    prednisolone_courses_df = build_event_table(
        df, ['Prednisolone_Name_Dosage_Quantity', 'Prednisolone_Dose', 'Prednisolone_Issue_Date'])
    antibiotic_courses_df = build_event_table(
        df, ['Antibiotic_Name_Dosage_Quantity', 'Antibiotic_Dose', 'Antibiotic_Issue_Date'])
    consultations_df = build_event_table(
        df, ['Consultation_Type', 'Consultation_Date'])
    referrals_df = build_event_table(
        df, ['Referral_to_Pulmonary_Rehab', 'Referral_to_Pulmonary_Rehab_Date'])
    influenza_df = build_event_table(df, ['Influenza_Vaccine', 'Influenza_Vaccine_Date'])
    pneumococcal_df = build_event_table(df, ['Pneumococcal_Vaccine', 'Pneumococcal_Vaccine_Date'])
    covid_df = build_event_table(df, ['COVID_Vaccine', 'COVID_Vaccine_Date'])

    # ---------- 5.5 Prescription/course counts onto patients_df ----------
    count_specs = [
        (inhaler_prescriptions_df, 'Inhaler_Prescription_Count'),
        (prednisolone_courses_df, 'Prednisolone_Course_Count'),
        (antibiotic_courses_df, 'Antibiotic_Course_Count'),
        (consultations_df, 'Consultation_Count'),
    ]
    for event_table, count_col in count_specs:
        counts = event_table.groupby('Patient_ID').size().rename(count_col)
        patients_df = patients_df.merge(counts, on='Patient_ID', how='left')
    count_cols = [c for _, c in count_specs]
    patients_df[count_cols] = patients_df[count_cols].fillna(0).astype(int)

    # ---------- 5.6 Evidence of Rescue Pack (text mention OR same-day pred+antibiotic) ----------
    RESCUE_PACK_COL_PAIRS = [
        ('Inhaler_Dose', 'Inhaler_Issue_Date'),
        ('Prednisolone_Dose', 'Prednisolone_Issue_Date'),
        ('Antibiotic_Dose', 'Antibiotic_Issue_Date'),
    ]

    # Signal 1: explicit "rescue pack" text mention, paired with that row's (parsed) date
    text_mentioned = pd.Series(False, index=df.index)
    text_mention_date = pd.Series(pd.NaT, index=df.index)
    for dose_col, date_col in RESCUE_PACK_COL_PAIRS:
        match = df[dose_col].str.contains('rescue pack', case=False, na=False)
        text_mentioned |= match
        text_mention_date = text_mention_date.mask(match, df[date_col])

    text_events = set(
        zip(df.loc[text_mentioned, 'Patient_ID'], text_mention_date[text_mentioned])
    )

    # Signal 2: prednisolone + antibiotic issued on the same (parsed) date
    prednisolone_dates = set(
        zip(df.loc[df['Prednisolone_Issue_Date'].notna(), 'Patient_ID'],
            df.loc[df['Prednisolone_Issue_Date'].notna(), 'Prednisolone_Issue_Date'])
    )
    antibiotic_dates = set(
        zip(df.loc[df['Antibiotic_Issue_Date'].notna(), 'Patient_ID'],
            df.loc[df['Antibiotic_Issue_Date'].notna(), 'Antibiotic_Issue_Date'])
    )
    inferred_events = prednisolone_dates & antibiotic_dates

    # Union: a (Patient_ID, date) counts once even if flagged by both signals
    rescue_pack_events = text_events | inferred_events

    rescue_pack_counts = (
        pd.Series([pid for pid, _ in rescue_pack_events], name='Patient_ID')
        .value_counts()
        .rename('Evidence_of_Rescue_Pack')
    )
    
    patients_df = patients_df.merge(rescue_pack_counts, on='Patient_ID', how='left')
    patients_df['Evidence_of_Rescue_Pack'] = patients_df['Evidence_of_Rescue_Pack'].fillna(0).astype(int)

    # ---------- 5.7 Vaccinated (Y/N) flags onto patients_df ----------
    vaccination_specs = [
        (influenza_df, 'Influenza_Vaccinated'),
        (pneumococcal_df, 'Pneumococcal_Vaccinated'),
        (covid_df, 'COVID_Vaccinated'),
    ]
    for event_table, flag_col in vaccination_specs:
        vaccinated_ids = set(event_table['Patient_ID'])
        patients_df[flag_col] = patients_df['Patient_ID'].isin(vaccinated_ids).map({True: 'Yes', False: pd.NA})

    # ---------- 6. Parse dates in every table ----------
    gp_tables = {
        'patients': patients_df,
        'inhaler_prescriptions': inhaler_prescriptions_df,
        'prednisolone_courses': prednisolone_courses_df,
        'antibiotic_courses': antibiotic_courses_df,
        'consultations': consultations_df,
        'referrals_pulmonary_rehab': referrals_df,
        'influenza_vaccinations': influenza_df,
        'pneumococcal_vaccinations': pneumococcal_df,
        'covid_vaccinations': covid_df,
    }
    for t in gp_tables.values():
        for c in t.columns:
            if 'Date' in c or c == 'DOB':
                t[c] = pd.to_datetime(t[c], dayfirst=True, errors='coerce')

    # ---------- 7. Standardise categorical text in patients_df ----------
    cat_cols = ['MRC_Code', 'COPD_Diagnosis', 'Depression', 'Anxiety']
    for c in cat_cols:
        patients_df[c] = patients_df[c].astype(str).str.strip().str.title().replace('Nan', pd.NA)
    patients_df['Postcode'] = patients_df['Postcode'].astype(str).str.strip().str.upper().replace('NAN', pd.NA)

    # ---------- 8. Sanity checks ----------
    if (patients_df['COPD_Diagnosis_Date'] < patients_df['DOB']).any():
        print(f"[{gp_id}] Data quality issue: some COPD diagnosis dates are before DOB")

    print(f"[{gp_id}] {len(patients_df)} patients, {len(df)} source rows")
    return gp_tables


# =========================================================================
# Excel export for analysts (NHS workers who work in Excel, not Python)
# One workbook, one sheet per table, plus a Summary index sheet up front.
# =========================================================================

HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
HEADER_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
BODY_FONT = Font(name='Arial', size=11)
TITLE_FONT = Font(name='Arial', bold=True, size=14)
DATE_FORMAT = 'DD/MM/YYYY'
MAX_COL_WIDTH = 45
MIN_COL_WIDTH = 10

# Excel sheet names are capped at 31 characters and can't repeat, so table
# names are mapped to short, analyst-friendly labels here.
SHEET_LABELS = {
    'patients': 'Patients',
    'inhaler_prescriptions': 'Inhaler Prescriptions',
    'prednisolone_courses': 'Prednisolone Courses',
    'antibiotic_courses': 'Antibiotic Courses',
    'consultations': 'Consultations',
    'referrals_pulmonary_rehab': 'Referrals - Pulm Rehab',
    'influenza_vaccinations': 'Influenza Vaccinations',
    'pneumococcal_vaccinations': 'Pneumococcal Vaccs',
    'covid_vaccinations': 'COVID Vaccinations',
}

SHEET_DESCRIPTIONS = {
    'patients': 'One row per patient: demographics, diagnosis, comorbidities and event counts.',
    'inhaler_prescriptions': 'One row per inhaler prescription issued.',
    'prednisolone_courses': 'One row per prednisolone course issued.',
    'antibiotic_courses': 'One row per antibiotic course issued.',
    'consultations': 'One row per recorded consultation.',
    'referrals_pulmonary_rehab': 'One row per pulmonary rehab referral.',
    'influenza_vaccinations': 'One row per influenza vaccination record.',
    'pneumococcal_vaccinations': 'One row per pneumococcal vaccination record.',
    'covid_vaccinations': 'One row per COVID vaccination record.',
}


def _write_table_sheet(wb, sheet_title, df):
    """Write one DataFrame to its own formatted, filterable sheet."""
    ws = wb.create_sheet(title=sheet_title[:31])

    # pandas NA/NaT types aren't writable by openpyxl - swap for plain None first.
    clean_df = df.astype(object).where(df.notna(), None)

    ws.append(list(clean_df.columns))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center')

    date_col_idxs = [i for i, c in enumerate(clean_df.columns, start=1) if 'Date' in c or c == 'DOB']

    for row in clean_df.itertuples(index=False, name=None):
        ws.append(list(row))

    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row_cells:
            cell.font = BODY_FONT
        for idx in date_col_idxs:
            row_cells[idx - 1].number_format = DATE_FORMAT

    for i, col in enumerate(clean_df.columns, start=1):
        sample = clean_df[col].map(lambda x: str(x) if pd.notna(x) else "").tolist()
        max_len = max([len(str(col))] + [len(v) for v in sample])
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(max_len + 2, MIN_COL_WIDTH),
            MAX_COL_WIDTH
        )
        
    ws.freeze_panes = 'A2'
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    return ws


def save_excel_workbook(tables, gp_sources, output_path):
    """Combine every cleaned table into one workbook: a Summary index sheet
    followed by one formatted, filterable sheet per table."""
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    summary = wb.create_sheet(title='Summary', index=0)
    summary.append(['COPD Evaluation - Cleaned Data Export'])
    summary['A1'].font = TITLE_FONT
    summary.append([f'Generated: {date.today().strftime("%d/%m/%Y")}'])
    summary['A2'].font = BODY_FONT
    summary.append([f'GP surgeries included: {", ".join(g["gp_id"] for g in gp_sources)}'])
    summary['A3'].font = BODY_FONT
    summary.append([])

    header_row = ['Sheet', 'Row Count', 'Description']
    summary.append(header_row)
    for cell in summary[summary.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for name, df in tables.items():
        label = SHEET_LABELS.get(name, name)
        summary.append([label, len(df), SHEET_DESCRIPTIONS.get(name, '')])
        for cell in summary[summary.max_row]:
            cell.font = BODY_FONT
        _write_table_sheet(wb, label, df)

    summary.column_dimensions['A'].width = 28
    summary.column_dimensions['B'].width = 12
    summary.column_dimensions['C'].width = 70
    summary.freeze_panes = 'A6'

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Excel workbook saved: {output_path} ({len(tables)} data sheets + Summary)")


# ---------- Run pipeline across all configured GP surgeries ----------
TABLE_NAMES = [
    'patients', 'inhaler_prescriptions', 'prednisolone_courses', 'antibiotic_courses',
    'consultations', 'referrals_pulmonary_rehab', 'influenza_vaccinations',
    'pneumococcal_vaccinations', 'covid_vaccinations',
]

per_gp_results = [process_gp_file(gp['gp_id'], gp['source_file']) for gp in GP_SOURCES]

tables = {
    name: pd.concat([result[name] for result in per_gp_results], ignore_index=True)
    for name in TABLE_NAMES
}
patients_df = tables['patients']

cross_gp_dupes = patients_df[patients_df.duplicated('Patient_ID', keep=False)]
if not cross_gp_dupes.empty:
    n_patients = cross_gp_dupes['Patient_ID'].nunique()
    gp_ids_involved = sorted(cross_gp_dupes['GP_ID'].unique())
    print(
        f"NOTE: {n_patients} Patient_ID(s) appear under more than one GP_ID "
        f"({gp_ids_involved}) - review before treating Patient_ID as a unique key."
    )

# ---------- 9. Save + summary ----------
# CSVs kept as-is for the pipeline/version control; the Excel workbook below
# is the analyst-facing deliverable (NHS workers who work in Excel, not Python).
for name, t in tables.items():
    t.to_csv(f'cleaned_data/combined_data/copd_{name}.csv', index=False)
    print(f"{name}: {t.shape}")

save_excel_workbook(
    tables,
    GP_SOURCES,
    'cleaned_data/combined_data/copd_combined_output.xlsx'
)

print(f"\nTotal GPs processed: {len(GP_SOURCES)}")
print(f"Total patients: {len(patients_df)}")
